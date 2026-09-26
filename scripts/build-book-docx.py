#!/usr/bin/env python
"""
Build the one Word document that holds the whole book transcription.

    python scripts/build-book-docx.py

The document is NOT edited in place and is never appended to by hand. It is
REBUILT, every time, from whatever per-page JSON transcripts exist under
extracted/transcripts/. Adding a section to the document therefore means
transcribing more pages, not touching the .docx -- which is what makes
"keep appending new sections to it" safe to do over many sittings: a rebuild
can never lose or duplicate what is already there, and the page order is
always the book's own.

Fonts: Devanagari and Tamil are set on the COMPLEX-SCRIPT slot (w:cs) as well
as the latin one. Word chooses the face for Devanagari from w:cs, so setting
only run.font.name leaves the text in a fallback face at the wrong size.
"""
import json, os, glob
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

ROOT  = r"G:\My Drive\Pooja Vidhi"
EXTR  = os.path.join(ROOT, "extracted")
TRANS = os.path.join(EXTR, "transcripts")
ILLUS = os.path.join(EXTR, "illustrations")
OUT   = os.path.join(ROOT, "Sampradaya Pooja Vidhi - Transcription.docx")

INDIC = "Nirmala UI"       # ships with Windows; covers Devanagari and Tamil
LATIN = "Calibri"
MONO  = "Consolas"

# The book's own order. `fallback_extent` is used until a section is transcribed.
SECTIONS = [
    {"slug": "0-general",        "title": "0. General Pooja Instructions",
     "pdf": "0.General Pooja instructions .pdf", "fallback_extent": "13 photographed pages"},
    {"slug": "1-common",         "title": "1. Common Initial Pooja (Purvanga)",
     "pdf": "1.Common initial Pooja .pdf", "fallback_extent": "26 photographed pages"},
    {"slug": "2-nitya",          "title": "2. Nitya (Daily) Pooja",
     "pdf": "2.Nitya Pooja .pdf", "fallback_extent": "26 photographed pages"},
    {"slug": "3-siddhivinayaka", "title": "3. Sri Siddhivinayaka Vratha Pooja",
     "pdf": "3.Siddhivinayak Vratha Pooja.pdf", "fallback_extent": "32 photographed pages"},
    {"slug": "4-varalakshmi",    "title": "4. Sri Varalakshmi Vrata Puja",
     "pdf": "4. Varalakshmi Vratha Pooja.pdf", "fallback_extent": "28 photographed pages"},
]

GREY   = RGBColor(0x55, 0x55, 0x55)
FAINT  = RGBColor(0x88, 0x88, 0x88)
RUST   = RGBColor(0x99, 0x33, 0x00)
GREEN  = RGBColor(0x00, 0x55, 0x22)
RED    = RGBColor(0xAA, 0x00, 0x00)
AMBER  = RGBColor(0x88, 0x66, 0x00)


def set_font(run, name, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rFonts.set(qn(slot), name)
    # The complex-script slot is ALWAYS the Indic face, even on a run whose
    # latin face is Calibri. Several English notes quote a Devanagari word
    # inline ("completed with पूजयामि"), and Word picks the face for those
    # characters from w:cs alone -- so a run set wholly to Calibri renders its
    # Devanagari in whatever fallback the machine happens to have.
    rFonts.set(qn("w:cs"), INDIC)
    if size is not None:
        run.font.size = Pt(size)
        # The complex-script size is a separate element. Without it, Devanagari
        # renders at the default size no matter what w:sz says.
        rPr.append(rPr.makeelement(qn("w:szCs"), {qn("w:val"): str(int(size * 2))}))
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def para(doc, text, font, size, **kw):
    p = doc.add_paragraph()
    set_font(p.add_run(text), font, size, **kw)
    return p


def load(slug):
    pages = []
    for path in sorted(glob.glob(os.path.join(TRANS, slug, "*.json"))):
        pages.append(json.load(open(path, encoding="utf-8")))
    return sorted(pages, key=lambda d: d["pdf_page"])


def render_block(doc, b, counts):
    t = b["type"]

    if t in ("h1", "h2", "h3"):
        size = {"h1": 17, "h2": 14, "h3": 12.5}[t]
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if t == "h3" else WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(" ".join(b["lines"])), INDIC, size, bold=True)

    elif t == "p":
        for line in b["lines"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_font(p.add_run(line), LATIN, 11)

    elif t == "list":
        for line in b["lines"]:
            p = doc.add_paragraph(style="List Bullet")
            set_font(p.add_run(line), LATIN, 11)

    elif t == "deva":
        for line in b["lines"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.2)
            set_font(p.add_run(line), INDIC, 13)
        counts["deva"] += len(b["lines"])

    elif t == "translit":
        for line in b["lines"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.2)
            set_font(p.add_run(line), LATIN, 11)
        counts["translit"] += len(b["lines"])

    elif t in ("note", "carryover"):
        for line in b["lines"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.left_indent = Inches(0.2)
            set_font(p.add_run(line), LATIN, 10, italic=True, color=GREY)

    elif t in ("warning", "finding"):
        label = "SOURCE DISCREPANCY" if t == "warning" else "FINDING"
        colour = RUST if t == "warning" else GREEN
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        set_font(p.add_run(label), LATIN, 9, bold=True, color=colour)
        for line in b["lines"]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            set_font(p.add_run(line), LATIN, 10, color=colour)
        counts[t] += 1

    elif t == "table":
        tbl = doc.add_table(rows=1, cols=len(b["columns"]))
        tbl.style = "Light Grid Accent 1"
        for i, c in enumerate(b["columns"]):
            set_font(tbl.rows[0].cells[i].paragraphs[0].add_run(c), LATIN, 9, bold=True)
        for row in b["rows"]:
            cells = tbl.add_row().cells
            for i, v in enumerate(row):
                # the last column is the English gloss; the rest are Devanagari
                last = (i == len(row) - 1)
                set_font(cells[i].paragraphs[0].add_run(v),
                         LATIN if last else INDIC,
                         9 if last else 11,
                         italic=last)
        doc.add_paragraph()

    elif t in ("namavali", "namavali_translit"):
        font = INDIC if t == "namavali" else LATIN
        size = 12 if t == "namavali" else 10.5
        tbl = doc.add_table(rows=0, cols=4)
        tbl.style = "Table Grid"
        # Without explicit widths Word gives the four columns equal share, and
        # the number columns -- which hold at most three digits -- steal room
        # from the names, wrapping the longer ones like lOkashOkavinaashinyai.
        tbl.autofit = False
        widths = [Inches(0.4), Inches(2.6), Inches(0.4), Inches(2.6)]
        n = b["start"]
        pair = []
        for name in b["names"]:
            pair.append((n, name))
            n += 1
            if len(pair) == 2:
                cells = tbl.add_row().cells
                for k, w in enumerate(widths):
                    cells[k].width = w
                for j, (num, nm) in enumerate(pair):
                    set_font(cells[j * 2].paragraphs[0].add_run(str(num)), LATIN, 8, color=FAINT)
                    set_font(cells[j * 2 + 1].paragraphs[0].add_run(nm), font, size)
                pair = []
        if pair:
            cells = tbl.add_row().cells
            for k, w in enumerate(widths):
                cells[k].width = w
            set_font(cells[0].paragraphs[0].add_run(str(pair[0][0])), LATIN, 8, color=FAINT)
            set_font(cells[1].paragraphs[0].add_run(pair[0][1]), font, size)
        doc.add_paragraph()
        counts["names"] += len(b["names"])

    elif t == "mudra":
        # The Purvanga file is largely a catalogue of mudras: a name, where on
        # the body it goes, and a drawing of it. Rendering those as three
        # separate blocks scatters each one down the page and loses which
        # drawing belongs to which name, so they are kept together in a row.
        tbl = doc.add_table(rows=1, cols=2)
        tbl.autofit = False
        left, right = tbl.rows[0].cells
        left.width, right.width = Inches(3.9), Inches(2.1)
        # Devanagari and transliteration on separate lines. Run together they
        # wrap mid-word ("sthaapitO / bhava"), which reads as a broken mantra.
        p = left.paragraphs[0]
        set_font(p.add_run(b["name_deva"]), INDIC, 14, bold=True)
        pt = left.add_paragraph()
        pt.paragraph_format.space_after = Pt(2)
        set_font(pt.add_run(b.get("name_translit", "")), LATIN, 11.5, bold=True)
        p2 = left.add_paragraph()
        set_font(p2.add_run(b.get("instruction", "")), LATIN, 10.5, italic=True, color=GREY)
        path = os.path.join(ILLUS, str(b.get("illustration")) + ".png")
        if b.get("illustration") and os.path.exists(path):
            rp = right.paragraphs[0]
            rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rp.add_run().add_picture(path, width=Inches(1.7))
            counts["illus"] += 1
        doc.add_paragraph()

    elif t == "illustration":
        path = os.path.join(ILLUS, b["id"] + ".png")
        if not os.path.exists(path):
            p = doc.add_paragraph()
            set_font(p.add_run("[illustration not yet cropped: " + b["id"] + "]"),
                     LATIN, 9, italic=True, color=RED)
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Inches(3.1))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(0)
        set_font(cap.add_run(b.get("caption", "")), LATIN, 9, italic=True, color=GREY)
        ref = doc.add_paragraph()
        ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(ref.add_run("extracted/illustrations/" + b["id"] + ".png"), MONO, 7.5, color=FAINT)
        counts["illus"] += 1


def main():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Sampradaya Vratha Pooja Vidhi"), LATIN, 26, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("A page-by-page transcription"), LATIN, 14, color=GREY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Published by Giri \u00b7 transcribed from photographs of the printed book"),
             LATIN, 10, italic=True, color=FAINT)
    doc.add_paragraph()

    counts = {"deva": 0, "translit": 0, "names": 0, "illus": 0, "warning": 0, "finding": 0}

    para(doc, "How to read this document", LATIN, 15, bold=True)
    for line in [
        "This is a transcription, not an edition. Everything below is set down as the book prints it, "
        "including its inconsistencies. Where the book contradicts itself the contradiction is recorded "
        "under a SOURCE DISCREPANCY heading rather than silently resolved, because deciding which "
        "reading is right is a question for a practitioner, not for a transcriber.",

        "Each printed page appears under its own page number, in the book's order. Mantras are given "
        "in Devanagari and then in the book's own transliteration, which is an ITRANS-like convention "
        "where capitals carry length and retroflexion (kSHeera, SHODasha, dEvi) and is NOT IAST. It has "
        "been copied exactly rather than converted, so that any line here can be checked against the page.",

        "Rubrics and stage directions, which the book sets in italic parentheses, are kept in italics.",

        "The illustrations are cropped from the same page photographs and are also saved as separate "
        "image files, named under each one, so they can be cleaned up and used elsewhere later.",
    ]:
        pp = doc.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pp.paragraph_format.space_after = Pt(8)
        set_font(pp.add_run(line), LATIN, 10.5)

    para(doc, "Sections", LATIN, 15, bold=True)
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    for i, h in enumerate(["Section", "Source file", "Extent", "Status"]):
        set_font(tbl.rows[0].cells[i].paragraphs[0].add_run(h), LATIN, 9, bold=True)

    present = {}
    for s in SECTIONS:
        pages = load(s["slug"])
        present[s["slug"]] = pages
        if pages:
            extent = "pp. %s-%s (%d pages)" % (pages[0]["book_page"], pages[-1]["book_page"], len(pages))
            bad = [pg for pg in pages if pg["quality"] == "damaged"]
            status = "Transcribed" if not bad else \
                "Transcribed; %d page(s) need re-photographing" % len(bad)
        else:
            extent, status = s["fallback_extent"], "Not yet transcribed"
        cells = tbl.add_row().cells
        for i, v in enumerate([s["title"], s["pdf"], extent, status]):
            set_font(cells[i].paragraphs[0].add_run(v), LATIN, 9)
    doc.add_paragraph()

    for s in SECTIONS:
        pages = present[s["slug"]]
        if not pages:
            continue
        doc.add_section(WD_SECTION.NEW_PAGE)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(16)
        set_font(p.add_run(s["title"]), LATIN, 20, bold=True)

        for pg in pages:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(2)
            set_font(p.add_run("Page %s" % pg["book_page"]), LATIN, 9, bold=True, color=FAINT)
            if pg["quality"] != "ok" and pg.get("quality_note"):
                q = doc.add_paragraph()
                q.paragraph_format.space_after = Pt(8)
                set_font(q.add_run(pg["quality_note"]), LATIN, 9, italic=True,
                         color=RED if pg["quality"] == "damaged" else AMBER)
            for b in pg["blocks"]:
                render_block(doc, b, counts)

    doc.save(OUT)
    print("wrote %s" % OUT)
    print("  %d Devanagari lines, %d transliteration lines, %d names" %
          (counts["deva"], counts["translit"], counts["names"]))
    print("  %d illustrations embedded, %d discrepancies, %d findings" %
          (counts["illus"], counts["warning"], counts["finding"]))


if __name__ == "__main__":
    main()
